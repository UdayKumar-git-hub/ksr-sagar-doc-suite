import docx
from docx.oxml import parse_xml
import json

def forensic_inspect(path):
    print("=" * 40, path, "=" * 40)
    doc = docx.Document(path)
    
    # Section
    sec = doc.sections[0]
    print(f"PAGE: width={sec.page_width.pt}pt ({sec.page_width.inches}in), height={sec.page_height.pt}pt ({sec.page_height.inches}in)")
    print(f"MARGINS: top={sec.top_margin.pt}pt, bottom={sec.bottom_margin.pt}pt, left={sec.left_margin.pt}pt, right={sec.right_margin.pt}pt")
    
    # Paragraphs before table
    print("\nPARAGRAPHS:")
    for i, p in enumerate(doc.paragraphs):
        pPr = p._p.pPr
        pPr_xml = pPr.xml if pPr is not None else ""
        print(f"P{i}: text='{p.text}'")
        print(f"    alignment={p.alignment}, space_before={p.paragraph_format.space_before}, space_after={p.paragraph_format.space_after}, line_spacing={p.paragraph_format.line_spacing}")
        for r_i, r in enumerate(p.runs):
            rPr = r._r.rPr
            rPr_xml = rPr.xml if rPr is not None else ""
            color = str(r.font.color.rgb) if (r.font.color and r.font.color.rgb) else "None"
            size = r.font.size.pt if r.font.size else "None"
            print(f"    Run {r_i}: text='{r.text}' font='{r.font.name}' size={size}pt bold={r.font.bold} italic={r.font.italic} color={color}")
    
    # Tables
    if doc.tables:
        t = doc.tables[0]
        print(f"\nTABLE: rows={len(t.rows)}, cols={len(t.columns)}")
        tblPr = t._tbl.tblPr
        print("tblPr xml:", tblPr.xml[:400] if tblPr is not None else "")
        tblGrid = t._tbl.tblGrid
        gridCols = [c.attrib for c in tblGrid] if tblGrid is not None else []
        print("tblGrid cols:", gridCols)
        
        # Check first 5 rows in detail
        for r_i in range(min(5, len(t.rows))):
            row = t.rows[r_i]
            trPr = row._tr.trPr
            trPr_xml = trPr.xml if trPr is not None else ""
            print(f"\n  Row {r_i} (cells={len(row.cells)}):")
            print(f"    trPr: {trPr_xml[:200]}")
            for c_i, c in enumerate(row.cells[:7]):
                tcPr = c._tc.tcPr
                tcPr_xml = tcPr.xml if tcPr is not None else ""
                # extract shd
                shd = [child.attrib for child in tcPr if child.tag.endswith('shd')] if tcPr is not None else []
                # extract tcBorders
                borders = [child.tag for child in tcPr if child.tag.endswith('tcBorders')] if tcPr is not None else []
                # extract tcMar
                mar = [child.attrib for child in tcPr if child.tag.endswith('tcMar')] if tcPr is not None else []
                # extract w
                w = [child.attrib for child in tcPr if child.tag.endswith('tcW')] if tcPr is not None else []
                # paragraph in cell
                cp = c.paragraphs[0] if c.paragraphs else None
                c_align = cp.alignment if cp else None
                c_run = cp.runs[0] if (cp and cp.runs) else None
                c_font = (c_run.font.name, c_run.font.size.pt if (c_run and c_run.font.size) else None, c_run.font.bold if c_run else None, str(c_run.font.color.rgb) if (c_run and c_run.font.color and c_run.font.color.rgb) else None) if c_run else None
                print(f"    Cell {c_i}: text='{c.text.strip().replace(chr(10), ' ')}' | w={w} | shd={shd} | align={c_align} | v_align={c.vertical_alignment} | font={c_font}")

forensic_inspect("templates/B.Tech Theory classes Day time table(Sat).docx")
forensic_inspect("templates/B.Tech Laboratories Day time table (Sat).docx")
