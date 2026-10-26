import docx
import os

doc_lab = docx.Document('templates/B.Tech Laboratories Day time table (Sat).docx')
print("Lab docx paragraphs:")
for i, p in enumerate(doc_lab.paragraphs):
    print(f"  P{i}: text='{p.text}'")

doc_th = docx.Document('templates/B.Tech Theory classes Day time table(Sat).docx')
print("Theory docx paragraphs:")
for i, p in enumerate(doc_th.paragraphs):
    print(f"  P{i}: text='{p.text}'")
