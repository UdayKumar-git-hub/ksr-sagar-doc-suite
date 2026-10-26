import docx

def compare_docs(ref_path, gen_path):
    print("=" * 60)
    print(f"COMPARING:\nREF: {ref_path}\nGEN: {gen_path}")
    print("=" * 60)
    
    doc_ref = docx.Document(ref_path)
    doc_gen = docx.Document(gen_path)
    
    # 1. Sections
    s_ref = doc_ref.sections[0]
    s_gen = doc_gen.sections[0]
    print(f"Margins REF: T={s_ref.top_margin.inches:.3f}, B={s_ref.bottom_margin.inches:.3f}, L={s_ref.left_margin.inches:.3f}, R={s_ref.right_margin.inches:.3f}")
    print(f"Margins GEN: T={s_gen.top_margin.inches:.3f}, B={s_gen.bottom_margin.inches:.3f}, L={s_gen.left_margin.inches:.3f}, R={s_gen.right_margin.inches:.3f}")
    
    # 2. Paragraphs before table
    print("\n--- PARAGRAPHS COMPARISON ---")
    for i in range(max(len(doc_ref.paragraphs), len(doc_gen.paragraphs))):
        p_ref = doc_ref.paragraphs[i] if i < len(doc_ref.paragraphs) else None
        p_gen = doc_gen.paragraphs[i] if i < len(doc_gen.paragraphs) else None
        
        if p_ref:
            r_ref = [(r.text[:20], r.font.name, r.font.size.pt if r.font.size else None, r.bold, str(r.font.color.rgb) if (r.font.color and r.font.color.rgb) else None) for r in p_ref.runs]
            print(f"REF P{i}: text='{p_ref.text[:30]}' align={p_ref.alignment} runs={r_ref}")
        if p_gen:
            r_gen = [(r.text[:20], r.font.name, r.font.size.pt if r.font.size else None, r.bold, str(r.font.color.rgb) if (r.font.color and r.font.color.rgb) else None) for r in p_gen.runs]
            print(f"GEN P{i}: text='{p_gen.text[:30]}' align={p_gen.alignment} runs={r_gen}")

    # 3. Table Rows
    t_ref = doc_ref.tables[0]
    t_gen = doc_gen.tables[0]
    print("\n--- TABLE HEADER ROW ---")
    print("REF cols count:", len(t_ref.columns), "Row 0 cells:", len(t_ref.rows[0].cells))
    print("GEN cols count:", len(t_gen.columns), "Row 0 cells:", len(t_gen.rows[0].cells))
    
    for c_i, cell in enumerate(t_ref.rows[0].cells):
        if c_i > 0 and cell._tc == t_ref.rows[0].cells[c_i-1]._tc: continue
        p = cell.paragraphs[0]
        r = p.runs[0] if p.runs else None
        print(f"  REF H-Col {c_i}: '{cell.text.strip()}' | align={p.alignment} | bold={r.bold if r else None} | font={r.font.name if r else None} | size={r.font.size.pt if (r and r.font.size) else None}")
        
    for c_i, cell in enumerate(t_gen.rows[0].cells):
        if c_i > 0 and cell._tc == t_gen.rows[0].cells[c_i-1]._tc: continue
        p = cell.paragraphs[0]
        r = p.runs[0] if p.runs else None
        print(f"  GEN H-Col {c_i}: '{cell.text.strip()}' | align={p.alignment} | bold={r.bold if r else None} | font={r.font.name if r else None} | size={r.font.size.pt if (r and r.font.size) else None}")

    print("\n--- TABLE SECTION ROW (Row 1) ---")
    c_ref = t_ref.rows[1].cells[0]
    p_ref = c_ref.paragraphs[0]
    r_ref = p_ref.runs[0] if p_ref.runs else None
    print(f"  REF Sec: '{c_ref.text.strip()}' | align={p_ref.alignment} | bold={r_ref.bold if r_ref else None} | size={r_ref.font.size.pt if (r_ref and r_ref.font.size) else None}")
    
    c_gen = t_gen.rows[1].cells[0]
    p_gen = c_gen.paragraphs[0]
    r_gen = p_gen.runs[0] if p_gen.runs else None
    print(f"  GEN Sec: '{c_gen.text.strip()}' | align={p_gen.alignment} | bold={r_gen.bold if r_gen else None} | size={r_gen.font.size.pt if (r_gen and r_gen.font.size) else None}")

    print("\n--- TABLE DATA ROW (Row 2) ---")
    for c_i, cell in enumerate(t_ref.rows[2].cells):
        if c_i > 0 and cell._tc == t_ref.rows[2].cells[c_i-1]._tc: continue
        p = cell.paragraphs[0]
        r = p.runs[0] if p.runs else None
        print(f"  REF Data {c_i}: '{cell.text.strip()}' | align={p.alignment} | bold={r.bold if r else None} | size={r.font.size.pt if (r and r.font.size) else None}")

    for c_i, cell in enumerate(t_gen.rows[2].cells):
        if c_i > 0 and cell._tc == t_gen.rows[2].cells[c_i-1]._tc: continue
        p = cell.paragraphs[0]
        r = p.runs[0] if p.runs else None
        print(f"  GEN Data {c_i}: '{cell.text.strip()}' | align={p.alignment} | bold={r.bold if r else None} | size={r.font.size.pt if (r and r.font.size) else None}")

compare_docs("templates/B.Tech Theory classes Day time table(Sat).docx", "outputs/test_theory.docx")
compare_docs("templates/B.Tech Laboratories Day time table (Sat).docx", "outputs/test_lab.docx")
