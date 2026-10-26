import docx

def compare_rows(ref_path, name):
    print("========================================")
    print(f"EXACT INSPECTION: {name} ({ref_path})")
    print("========================================")
    doc = docx.Document(ref_path)
    
    # Paragraphs
    print("--- PARAGRAPHS ---")
    for i, p in enumerate(doc.paragraphs):
        p_align = p.alignment
        runs_detail = []
        for r in p.runs:
            bold = r.bold
            font = r.font.name
            sz = r.font.size.pt if r.font.size else None
            col = str(r.font.color.rgb) if (r.font.color and r.font.color.rgb) else None
            runs_detail.append((r.text, font, sz, bold, col))
        print(f"P{i} [align={p_align}, sb={p.paragraph_format.space_before}, sa={p.paragraph_format.space_after}]: {runs_detail}")
        
    # Table
    if doc.tables:
        t = doc.tables[0]
        print(f"\n--- TABLE (rows={len(t.rows)}, cols={len(t.columns)}) ---")
        
        # Check first 8 rows
        for r_idx in range(min(8, len(t.rows))):
            row = t.rows[r_idx]
            print(f"\nRow {r_idx} (cells={len(row.cells)}):")
            for c_idx, cell in enumerate(row.cells):
                # check if cell is duplicate due to merge
                if c_idx > 0 and cell._tc == row.cells[c_idx - 1]._tc:
                    continue
                p = cell.paragraphs[0] if cell.paragraphs else None
                p_align = p.alignment if p else None
                runs_info = []
                if p:
                    for r in p.runs:
                        runs_info.append({
                            'text': r.text,
                            'bold': r.bold,
                            'font': r.font.name,
                            'size': r.font.size.pt if r.font.size else None,
                            'color': str(r.font.color.rgb) if (r.font.color and r.font.color.rgb) else None
                        })
                tcPr = cell._tc.tcPr
                shd = [child.attrib for child in tcPr if child.tag.endswith('shd')] if tcPr is not None else []
                v_align = cell.vertical_alignment
                print(f"  Col {c_idx} [align={p_align}, v_align={v_align}, shd={shd}]: {runs_info}")

compare_rows("templates/B.Tech Theory classes Day time table(Sat).docx", "REFERENCE THEORY")
compare_rows("templates/B.Tech Laboratories Day time table (Sat).docx", "REFERENCE LAB")
