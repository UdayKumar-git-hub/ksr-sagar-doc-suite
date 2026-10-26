import os
import docx

def inspect(path):
    print('=== ' + path + ' ===')
    doc = docx.Document(path)
    print('Paragraphs:', len(doc.paragraphs))
    for i, p in enumerate(doc.paragraphs[:10]):
        if p.text.strip():
            runs = [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.font.bold, str(r.font.color.rgb) if (r.font.color and r.font.color.rgb) else None) for r in p.runs]
            print(f'P{i}: text="{p.text}", align={p.alignment}, runs={runs}')
    print('Tables:', len(doc.tables))
    if doc.tables:
        t = doc.tables[0]
        print(f'Table: rows={len(t.rows)}, cols={len(t.columns)}')
        for r_idx in range(min(6, len(t.rows))):
            row = t.rows[r_idx]
            vals = [c.text.replace('\n', ' ') for c in row.cells]
            print(f'  Row {r_idx}: {vals[:7]}')
            c0 = row.cells[0]
            tcPr = c0._tc.get_or_add_tcPr()
            shd = [child.attrib for child in tcPr if child.tag.endswith('shd')]
            fnt = None
            if c0.paragraphs and c0.paragraphs[0].runs:
                r0 = c0.paragraphs[0].runs[0]
                fnt = (r0.font.name, r0.font.size.pt if r0.font.size else None, r0.font.bold, str(r0.font.color.rgb) if (r0.font.color and r0.font.color.rgb) else None)
            print(f'    shd={shd}, font={fnt}')

inspect(r'templates/B.Tech Laboratories Day time table (Sat).docx')
inspect(r'templates/B.Tech Theory classes Day time table(Sat).docx')
