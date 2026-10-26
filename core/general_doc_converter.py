import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import pdfplumber

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_borders(cell, sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="CCCCCC"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="CCCCCC"/>
            <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="CCCCCC"/>
            <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="CCCCCC"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def convert_general_pdf_to_faculty_docx(pdf_path, output_docx_path, doc_title="Faculty Institutional Document", doc_type="Academic Document"):
    """
    Extracts text and tables from any uploaded PDF and renders it into the KSR Sagar official Word template format.
    """
    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11.0)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

    # Institutional Header Block
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("KSR SAGAR GROUP OF INSTITUTIONS")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(16)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(234, 88, 12) # Warm Saffron Orange

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Academic Planning, Monitoring & Continuing Studies (PMCS)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Document Title Block
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_title = p_title.add_run(doc_title)
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    # Divider bar
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(12)
    r_div = p_div.add_run("―" * 58)
    r_div.font.color.rgb = RGBColor(203, 213, 225)

    extracted_pages = []

    with pdfplumber.open(pdf_path) as pdf:
        for p_idx, page in enumerate(pdf.pages):
            p_num = p_idx + 1
            tables = page.extract_tables()
            text = page.extract_text() or ""
            
            # Heading for page if multi-page
            if len(pdf.pages) > 1:
                p_sec = doc.add_paragraph()
                p_sec.paragraph_format.space_before = Pt(12)
                p_sec.paragraph_format.space_after = Pt(4)
                r_sec = p_sec.add_run(f"Section {p_num}")
                r_sec.font.name = "Calibri"
                r_sec.font.size = Pt(12)
                r_sec.font.bold = True
                r_sec.font.color.rgb = RGBColor(234, 88, 12)

            # Insert tables if detected
            if tables:
                for tbl_data in tables:
                    if not tbl_data or len(tbl_data) == 0:
                        continue
                    num_rows = len(tbl_data)
                    num_cols = max(len(r) for r in tbl_data if r)
                    if num_cols == 0:
                        continue
                        
                    t = doc.add_table(rows=num_rows, cols=num_cols)
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER
                    
                    for r_i, row in enumerate(tbl_data):
                        for c_i, val in enumerate(row):
                            if c_i < num_cols:
                                cell = t.cell(r_i, c_i)
                                cell.text = str(val or "").strip()
                                set_cell_borders(cell, sz="4")
                                p_c = cell.paragraphs[0]
                                p_c.paragraph_format.space_before = Pt(2)
                                p_c.paragraph_format.space_after = Pt(2)
                                if r_i == 0:
                                    set_cell_background(cell, "E0E0E0")
                                    p_c.runs[0].font.bold = True if p_c.runs else None
                                    p_c.runs[0].font.size = Pt(9) if p_c.runs else None
                                else:
                                    if p_c.runs:
                                        p_c.runs[0].font.size = Pt(8.5)
                    doc.add_paragraph().paragraph_format.space_after = Pt(6)

            # Add body text paragraphs
            lines = [l.strip() for l in text.split("\n") if l.strip()]
            for line in lines:
                # Skip duplicate timetable lines if already in tables
                if tables and len(tables) > 0 and len(line) < 30 and ("CLASS/PERIOD" in line or "Room :" in line):
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(line)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(30, 41, 59)

            extracted_pages.append({
                "page": p_num,
                "text_length": len(text),
                "tables_count": len(tables)
            })

    # Official signature block
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(36)
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sig1 = p_sig.add_run("Authorized Signatory\n")
    r_sig1.font.bold = True
    r_sig1.font.size = Pt(10.5)
    r_sig2 = p_sig.add_run("Office of the Dean, Academic Planning\nKSR Sagar Institution")
    r_sig2.font.size = Pt(9.5)
    r_sig2.font.color.rgb = RGBColor(71, 85, 105)

    doc.save(output_docx_path)
    return output_docx_path, extracted_pages
