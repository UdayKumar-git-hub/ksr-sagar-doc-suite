import docx
from docx.shared import Inches
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
col_widths = [Inches(0.56), Inches(1.19), Inches(1.81), Inches(0.81), Inches(1.44), Inches(0.75), Inches(0.88)]
table = doc.add_table(rows=1, cols=7)
table.autofit = False
tblGrid = table._tbl.tblGrid
for child in list(tblGrid):
    tblGrid.remove(child)
for w in col_widths:
    dxa = int(w.inches * 1440)
    tblGrid.append(parse_xml(f'<w:gridCol {nsdecls("w")} w:w="{dxa}"/>'))
for i, col in enumerate(table.columns):
    col.width = col_widths[i]
for i, cell in enumerate(table.rows[0].cells):
    cell.width = col_widths[i]

doc.save('outputs/test_grid.docx')
doc_chk = docx.Document('outputs/test_grid.docx')
for i, c in enumerate(doc_chk.tables[0].columns):
    print(f'Col {i} width: {c.width.inches:.3f}')
