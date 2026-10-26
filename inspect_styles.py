import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

def inspect_doc(doc_path):
    doc = docx.Document(doc_path)
    res = {}
    
    # Section properties
    sec = doc.sections[0]
    res['page_width_in'] = sec.page_width.inches
    res['page_height_in'] = sec.page_height.inches
    res['orientation'] = str(sec.orientation)
    res['margins'] = {
        'top_in': sec.top_margin.inches,
        'bottom_in': sec.bottom_margin.inches,
        'left_in': sec.left_margin.inches,
        'right_in': sec.right_margin.inches,
    }
    
    # Paragraphs info
    res['paragraphs'] = []
    for p in doc.paragraphs:
        runs_info = [{'text': r.text, 'bold': r.bold, 'size': r.font.size.pt if r.font.size else None, 'font': r.font.name} for r in p.runs]
        res['paragraphs'].append({
            'text': p.text,
            'align': str(p.alignment),
            'runs': runs_info
        })
        
    # Table info
    t = doc.tables[0]
    res['table_cols'] = len(t.columns)
    res['table_rows'] = len(t.rows)
    
    # Inspect sample rows
    rows_sample = []
    for r_idx in range(min(15, len(t.rows))):
        row = t.rows[r_idx]
        cells_data = []
        for c in row.cells:
            p_first = c.paragraphs[0] if c.paragraphs else None
            font_name = p_first.runs[0].font.name if (p_first and p_first.runs and p_first.runs[0].font) else None
            font_size = p_first.runs[0].font.size.pt if (p_first and p_first.runs and p_first.runs[0].font and p_first.runs[0].font.size) else None
            is_bold = p_first.runs[0].bold if (p_first and p_first.runs) else None
            cells_data.append({
                'text': c.text.strip().replace('\n', ' '),
                'font': font_name,
                'size': font_size,
                'bold': is_bold
            })
        rows_sample.append(cells_data)
    res['rows_sample'] = rows_sample
    
    # Column widths
    col_widths = [col.width.inches if col.width else None for col in t.columns]
    res['col_widths'] = col_widths
    
    return res

lab_info = inspect_doc(r'C:\Users\ADMIN\.gemini\antigravity-ide\scratch\ksr-sagar-doc-suite\templates\B.Tech Laboratories Day time table (Sat).docx')
theory_info = inspect_doc(r'C:\Users\ADMIN\.gemini\antigravity-ide\scratch\ksr-sagar-doc-suite\templates\B.Tech Theory classes Day time table(Sat).docx')

print('=== LAB INFO ===')
print('Page size:', lab_info['page_width_in'], 'x', lab_info['page_height_in'], 'Margins:', lab_info['margins'])
print('Paragraphs:', [p['text'] for p in lab_info['paragraphs']])
print('Col widths:', lab_info['col_widths'])
for i, r in enumerate(lab_info['rows_sample'][:6]):
    texts = [c['text'] for c in r]
    print(f'Row {i}: {texts}')

print('\n=== THEORY INFO ===')
print('Page size:', theory_info['page_width_in'], 'x', theory_info['page_height_in'], 'Margins:', theory_info['margins'])
print('Paragraphs:', [p['text'] for p in theory_info['paragraphs']])
print('Col widths:', theory_info['col_widths'])
for i, r in enumerate(theory_info['rows_sample'][:6]):
    texts = [c['text'] for c in r]
    print(f'Row {i}: {texts}')
