import docx

doc = docx.Document('templates/B.Tech Theory classes Day time table(Sat).docx')
print('=== THEORY PARAGRAPHS ===')
for i, p in enumerate(doc.paragraphs):
    print(f'P{i}: text="{p.text}", align={p.alignment}')
    for r_i, r in enumerate(p.runs):
        print(f'   R{r_i}: font={r.font.name}, size={r.font.size.pt if r.font.size else None}, bold={r.font.bold}, color={r.font.color.rgb if r.font.color else None}')

t = doc.tables[0]
print('\n=== THEORY TABLE ROW 0 (Headers) ===')
for c_i, c in enumerate(t.rows[0].cells):
    cp = c.paragraphs[0] if c.paragraphs else None
    run = cp.runs[0] if (cp and cp.runs) else None
    txt = c.text.strip().replace('\n', ' ')
    print(f'Col {c_i}: text="{txt}", align={cp.alignment if cp else None}, size={run.font.size.pt if (run and run.font.size) else None}, bold={run.font.bold if run else None}')

print('\n=== THEORY TABLE ROW 1 (Section) ===')
c = t.rows[1].cells[0]
cp = c.paragraphs[0]
run = cp.runs[0] if cp.runs else None
print(f'Section: text="{c.text.strip()}", align={cp.alignment}, bold={run.font.bold if run else None}, size={run.font.size.pt if (run and run.font.size) else None}')

print('\n=== THEORY TABLE ROW 2 (Data) ===')
for c_i, c in enumerate(t.rows[2].cells):
    cp = c.paragraphs[0] if c.paragraphs else None
    run = cp.runs[0] if (cp and cp.runs) else None
    txt = c.text.strip().replace('\n', ' ')
    print(f'Col {c_i}: text="{txt}", align={cp.alignment if cp else None}, size={run.font.size.pt if (run and run.font.size) else None}')
