import docx
import os
import fitz

print('=== INSPECTING PDF ===')
pdf_path = r'c:\Users\ADMIN\Downloads\Telegram Desktop\IARE - Day Time Table - DEAN_PG.pdf'
doc_pdf = fitz.open(pdf_path)
print('PDF pages count:', len(doc_pdf))
for pno in range(min(5, len(doc_pdf))):
    p = doc_pdf[pno]
    rect = p.rect
    orient = 'landscape' if rect.width > rect.height else 'portrait'
    print(f'Page {pno+1}: width={rect.width}, height={rect.height}, orientation={orient}')
    text = p.get_text()
    lines = [l.strip() for l in text.split('\n') if l.strip()]
    print(f'Page {pno+1} first 8 lines: {lines[:8]}')

print('\n=== INSPECTING B.Tech Laboratories Day time table (Sat).docx ===')
lab_docx = docx.Document(r'C:\Users\ADMIN\.gemini\antigravity-ide\scratch\ksr-sagar-doc-suite\templates\B.Tech Laboratories Day time table (Sat).docx')
print('Lab Doc paragraphs:', len(lab_docx.paragraphs))
for i, p in enumerate(lab_docx.paragraphs[:5]):
    if p.text:
        print(f'P{i}: {p.text}')
print('Lab Doc tables:', len(lab_docx.tables))
for t_idx, t in enumerate(lab_docx.tables):
    print(f'Table {t_idx}: rows={len(t.rows)}, cols={len(t.columns)}')
    for r_idx in range(min(4, len(t.rows))):
        row_txt = [c.text.replace('\n', ' ')[:25] for c in t.rows[r_idx].cells]
        print(f'  Row {r_idx}: {row_txt[:7]}')

print('\n=== INSPECTING B.Tech Theory classes Day time table(Sat).docx ===')
theory_docx = docx.Document(r'C:\Users\ADMIN\.gemini\antigravity-ide\scratch\ksr-sagar-doc-suite\templates\B.Tech Theory classes Day time table(Sat).docx')
print('Theory Doc paragraphs:', len(theory_docx.paragraphs))
for i, p in enumerate(theory_docx.paragraphs[:5]):
    if p.text:
        print(f'P{i}: {p.text}')
print('Theory Doc tables:', len(theory_docx.tables))
for t_idx, t in enumerate(theory_docx.tables):
    print(f'Table {t_idx}: rows={len(t.rows)}, cols={len(t.columns)}')
    for r_idx in range(min(4, len(t.rows))):
        row_txt = [c.text.replace('\n', ' ')[:25] for c in t.rows[r_idx].cells]
        print(f'  Row {r_idx}: {row_txt[:7]}')
