import docx
import os

def extract_images_and_paras(path, prefix):
    doc = docx.Document(path)
    os.makedirs('extracted_assets', exist_ok=True)
    
    for i, p in enumerate(doc.paragraphs):
        print(f"[{prefix} P{i}] text='{p.text}' align={p.alignment}")
        # check runs for drawing/blip
        for r_i, r in enumerate(p.runs):
            # check oxml
            xml = r._r.xml
            if 'w:drawing' in xml or 'blip' in xml or 'graphic' in xml:
                print(f"  --> Run {r_i} has DRAWING / IMAGE!")
                
    # Also check tables for images
    for t_i, t in enumerate(doc.tables):
        for r_i, r in enumerate(t.rows[:2]):
            for c_i, c in enumerate(r.cells):
                for p_i, p in enumerate(c.paragraphs):
                    for run in p.runs:
                        if 'w:drawing' in run._r.xml or 'blip' in run._r.xml:
                            print(f"  --> Table {t_i} Row {r_i} Col {c_i} has DRAWING!")

    part = doc.part
    img_idx = 0
    for rel_id, rel in part.related_parts.items():
        if "image" in rel.content_type:
            img_idx += 1
            ext = rel.content_type.split('/')[-1]
            out_img = os.path.join('extracted_assets', f'{prefix}_logo_{img_idx}.{ext}')
            with open(out_img, 'wb') as f:
                f.write(rel.blob)
            print(f"Extracted image {rel_id} to {out_img} ({len(rel.blob)} bytes)")

extract_images_and_paras(r"templates/B.Tech Laboratories Day time table (Sat).docx", "lab")
extract_images_and_paras(r"templates/B.Tech Theory classes Day time table(Sat).docx", "theory")
